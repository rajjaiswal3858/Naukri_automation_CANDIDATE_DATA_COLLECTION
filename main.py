import argparse
import os
import random
import subprocess
import hashlib
import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlparse

import pymupdf as fitz
import pandas as pd
from docx import Document
from playwright.sync_api import sync_playwright


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

DATA_DIR = BASE_DIR / "local_data"
CANDIDATES_DIR = DATA_DIR / "candidates"
FAILED_DIR = DATA_DIR / "failed"
EXPORT_DIR = DATA_DIR / "exports"
LOG_DIR = DATA_DIR / "logs"

SESSION_DIR = BASE_DIR / "session"

# Dedicated persistent Chrome profile used only for this automation.
#
# The Naukri Launcher opens the dashboard with `start chrome "<url>"`,
# which always uses Chrome's default profile directory. The per-user
# Chrome policy UserDataDir (set by ensure_chrome_policy) makes that
# default directory this folder, so the Launcher's tab lands in the
# Chrome instance the script controls over CDP.
CHROME_PROFILE_DIR = BASE_DIR / "chrome_naukri_profile"

CHROME_POLICY_KEY = r"Software\Policies\Google\Chrome"

CDP_URL = "http://127.0.0.1:9222"

# How long to wait for the user to log in through the Launcher.
LOGIN_WAIT_SECONDS = 600

# Random pause (min, max seconds) after opening each candidate profile.
PROFILE_DELAY_SECONDS = (4, 9)

EXCEL_FILE = EXPORT_DIR / "candidates.xlsx"
PROGRESS_FILE = DATA_DIR / "progress.json"

DAILY_LIMIT = 500

for directory in [
    CANDIDATES_DIR,
    FAILED_DIR,
    EXPORT_DIR,
    LOG_DIR,
    SESSION_DIR,
    CHROME_PROFILE_DIR,
]:
    directory.mkdir(parents=True, exist_ok=True)


# ============================================================
# LOGGING
# ============================================================

def log(message, level="INFO"):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{timestamp}] {level}: {message}"

    print(line)

    log_file = LOG_DIR / f"{datetime.now():%Y-%m-%d}.log"

    with open(log_file, "a", encoding="utf-8") as f:
        f.write(line + "\n")


# ============================================================
# HELPERS
# ============================================================

def safe_filename(value):
    value = value or "unknown"

    value = re.sub(r'[<>:"/\\|?*]', "_", value)
    value = re.sub(r"\s+", "_", value)

    return value[:100]


def sha256_file(path):
    sha = hashlib.sha256()

    with open(path, "rb") as f:
        while True:
            chunk = f.read(1024 * 1024)

            if not chunk:
                break

            sha.update(chunk)

    return sha.hexdigest()


def load_progress():
    if not PROGRESS_FILE.exists():
        return {
            "processed": 0,
            "successful": 0,
            "failed": 0,
            "duplicates": 0,
        }

    try:
        return json.loads(
            PROGRESS_FILE.read_text(
                encoding="utf-8"
            )
        )

    except Exception:
        return {
            "processed": 0,
            "successful": 0,
            "failed": 0,
            "duplicates": 0,
        }


def save_progress(progress):
    PROGRESS_FILE.write_text(
        json.dumps(
            progress,
            indent=2
        ),
        encoding="utf-8",
    )


def is_naukri_url(url):
    return "naukri.com" in (url or "").lower()


# ============================================================
# RESUME PARSER
# ============================================================

def extract_pdf(path):
    text = []

    document = fitz.open(path)

    try:
        for page in document:
            text.append(page.get_text())

    finally:
        document.close()

    return "\n".join(text)


def extract_docx(path):
    document = Document(path)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(paragraphs)


def extract_doc(path):
    """
    Pull readable text out of an old binary Word (.doc) file.

    No Word install is needed: the text is stored as plain 8-bit or
    UTF-16 runs inside the file, which is enough to find the name,
    email and phone number.
    """

    data = Path(path).read_bytes()

    runs = re.findall(rb"[\x20-\x7e\r\n\t]{4,}", data)

    runs += [
        match.decode("utf-16-le").encode("ascii", "ignore")
        for match in re.findall(rb"(?:[\x20-\x7e\r\n\t]\x00){4,}", data)
    ]

    return "\n".join(
        run.decode("ascii", "ignore").strip()
        for run in runs
        if run.strip()
    )


def parse_resume(path):
    extension = path.suffix.lower()

    if extension == ".pdf":
        text = extract_pdf(path)

    elif extension == ".docx":
        text = extract_docx(path)

    elif extension == ".doc":
        # The first readable run in a .doc is formatting data, not the
        # candidate's name, so only email/phone are taken from it.
        fields = extract_resume_fields(extract_doc(path))
        fields["parsed_name"] = ""
        return fields

    else:
        raise ValueError(
            f"Unsupported resume format: {extension}"
        )

    return extract_resume_fields(text)


def extract_resume_fields(text):
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    email = ""

    email_match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text,
    )

    if email_match:
        email = email_match.group(0)

    phone = ""

    phone_match = re.search(
        r"(?:\+91[\s-]?)?[6-9]\d{9}",
        text,
    )

    if phone_match:
        phone = phone_match.group(0)

    name = lines[0] if lines else ""

    return {
        "resume_text": text,
        "parsed_name": name,
        "parsed_email": email,
        "parsed_phone": phone,
    }


# ============================================================
# DUPLICATE CHECK
# ============================================================

def find_duplicate(resume_hash):
    if not CANDIDATES_DIR.exists():
        return None

    for candidate_dir in CANDIDATES_DIR.iterdir():

        if not candidate_dir.is_dir():
            continue

        candidate_json = candidate_dir / "candidate.json"

        if not candidate_json.exists():
            continue

        try:
            data = json.loads(
                candidate_json.read_text(
                    encoding="utf-8"
                )
            )

            if data.get("resume_hash") == resume_hash:
                return data

        except Exception:
            continue

    return None


# ============================================================
# SAVE CANDIDATE
# ============================================================

def save_candidate(candidate, resume_path):

    candidate_id = safe_filename(
        candidate.get("candidate_id", "unknown")
    )

    candidate_name = safe_filename(
        candidate.get("name", "unknown")
    )

    folder = (
        CANDIDATES_DIR
        / f"{candidate_id}_{candidate_name}"
    )

    folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    final_resume = folder / resume_path.name

    # If the destination already exists, remove it first.
    if final_resume.exists():
        final_resume.unlink()

    resume_path.replace(final_resume)

    resume_hash = sha256_file(final_resume)

    parsed = parse_resume(final_resume)

    candidate.update(parsed)

    candidate["resume_file"] = final_resume.name
    candidate["resume_hash"] = resume_hash
    candidate["downloaded_at"] = datetime.now().isoformat()
    candidate["parsed_at"] = datetime.now().isoformat()
    candidate["status"] = "success"

    json_file = folder / "candidate.json"

    json_file.write_text(
        json.dumps(
            candidate,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    return candidate


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_excel():

    rows = []

    if not CANDIDATES_DIR.exists():
        log("No candidate directory found.")

        return

    for candidate_dir in CANDIDATES_DIR.iterdir():

        if not candidate_dir.is_dir():
            continue

        json_file = candidate_dir / "candidate.json"

        if not json_file.exists():
            continue

        try:
            data = json.loads(
                json_file.read_text(
                    encoding="utf-8"
                )
            )

            # Resume text can be very large.
            data.pop("resume_text", None)

            rows.append(data)

        except Exception as e:
            log(
                f"Could not read {json_file}: {e}",
                "WARNING",
            )

    if not rows:
        log("No candidates available for Excel export.")

        return

    df = pd.DataFrame(rows)

    # Convert list values into Excel-friendly strings.
    for column in df.columns:
        df[column] = df[column].apply(
            lambda value:
            ", ".join(map(str, value))
            if isinstance(value, list)
            else value
        )

    df.to_excel(
        EXCEL_FILE,
        index=False,
    )

    log(
        f"Excel updated: {EXCEL_FILE}",
        "SUCCESS",
    )

    export_structured_excel()


# ============================================================
# STRUCTURED EXCEL EXPORT
# ============================================================

STRUCTURED_EXCEL_FILE = EXPORT_DIR / "candidates_structured.xlsx"

KNOWN_CITIES = [
    "Hyderabad", "Bengaluru", "Bangalore", "Mumbai", "Navi Mumbai",
    "Thane", "Pune", "Chennai", "Kolkata", "Delhi", "New Delhi",
    "Noida", "Greater Noida", "Gurgaon", "Gurugram", "Faridabad",
    "Ghaziabad", "Ahmedabad", "Vadodara", "Surat", "Jaipur",
    "Indore", "Bhopal", "Lucknow", "Kanpur", "Chandigarh", "Mohali",
    "Nagpur", "Nashik", "Coimbatore", "Kochi", "Cochin",
    "Thiruvananthapuram", "Trivandrum", "Visakhapatnam", "Vijayawada",
    "Mysore", "Mysuru", "Mangalore", "Bhubaneswar", "Patna", "Ranchi",
    "Dehradun", "Ludhiana", "Amritsar", "Raipur", "Guwahati", "Goa",
    "Madurai", "Tiruchirappalli", "Warangal", "Udaipur", "Jodhpur",
    "Kota", "Agra", "Varanasi", "Allahabad", "Prayagraj", "Meerut",
]

QUALIFICATION_PATTERNS = {
    "CA": r"\bCA\b(?!\s*Inter)|Chartered Accountant",
    "CA Inter": r"\bCA[\s-]*Inter(?:mediate)?\b",
    "CPA": r"\bCPA\b|Certified Public Accountant",
    "ACCA": r"\bACCA\b",
    "CMA": r"\bCMA\b|\bICWA\b|Cost (?:and|&) Management Accountant",
    "CS": r"\bCompany Secretary\b",
    "CFA": r"\bCFA\b",
    "EA": r"\bEnrolled Agent\b",
    "MBA": r"\bMBA\b|\bPGDM\b",
    "M.Com": r"\bM\.?\s?Com\b",
    "B.Com": r"\bB\.?\s?Com\b",
    "BBA": r"\bBBA\b",
    "B.Tech": r"\bB\.?\s?Tech\b|\bB\.?E\.\b",
    "Graduate": r"\bGraduat(?:e|ion)\b",
}

SKILL_KEYWORDS = [
    "US GAAP", "IFRS", "IND AS", "GAAP", "SOX",
    "Record to Report", "R2R", "Procure to Pay", "P2P",
    "Order to Cash", "O2C", "Accounts Payable", "Accounts Receivable",
    "General Ledger", "Bank Reconciliation", "Reconciliation",
    "Month End Close", "Year End Close", "Financial Reporting",
    "FP&A", "Budgeting", "Forecasting", "Audit", "Payroll",
    "Bookkeeping", "Sales Tax", "1099", "Tax Return", "Form 1040",
    "Form 1120", "Form 1065", "Fixed Assets", "Intercompany",
    "Accruals", "Revenue Recognition", "Variance Analysis",
    "QuickBooks", "Xero", "NetSuite", "SAP", "Oracle", "Sage",
    "Tally", "Yardi", "Bill.com", "Expensify", "BlackLine",
    "Dynamics", "Zoho Books", "FreshBooks", "Excel", "Power BI",
    "VLOOKUP", "Pivot",
]

DESIGNATION_WORDS = (
    "accountant", "accounting", "manager", "analyst", "executive",
    "associate", "consultant", "lead", "specialist", "officer",
    "controller", "auditor", "head", "director", "supervisor",
    "bookkeeper", "processor", "senior", "sr.", "assistant",
)


def structured_fields_from_resume(text):
    """
    Pull structured details out of a resume's plain text.
    Every field is best-effort and left blank when not found.
    """

    text = text or ""
    flat = re.sub(r"\s+", " ", text)

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    # Phone: Indian mobiles, with or without +91 / spaces / dashes.
    phone = ""

    phone_match = re.search(
        r"(?:\+?91[\s-]*)?(?<!\d)[6-9]\d{2}[\s-]?\d{3}[\s-]?\d{4}(?!\d)",
        flat,
    )

    if phone_match:
        phone = re.sub(r"[^\d+]", "", phone_match.group(0))

    linkedin = ""

    linkedin_match = re.search(
        r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[A-Za-z0-9_%-]+/?",
        flat,
        re.IGNORECASE,
    )

    if linkedin_match:
        linkedin = linkedin_match.group(0)

    # Total experience: largest "N years" figure that reads like a
    # career total (ignore silly values like 1990 years).
    experience = ""

    years = [
        float(match)
        for match in re.findall(
            r"(\d{1,2}(?:\.\d{1,2})?)\s*\+?\s*(?:years|yrs|year)",
            flat,
            re.IGNORECASE,
        )
    ]

    years = [value for value in years if 0 < value <= 45]

    if years:
        best = max(years)
        experience = int(best) if best.is_integer() else best

    location = ""

    for city in KNOWN_CITIES:
        if re.search(rf"\b{re.escape(city)}\b", flat, re.IGNORECASE):
            location = city
            break

    qualifications = [
        label
        for label, pattern in QUALIFICATION_PATTERNS.items()
        if re.search(pattern, flat)
    ]

    skills = [
        skill
        for skill in SKILL_KEYWORDS
        if re.search(
            rf"(?<![A-Za-z]){re.escape(skill)}(?![A-Za-z])",
            flat,
            re.IGNORECASE,
        )
    ]

    # Designation: first short line near the top that looks like a
    # job title.
    designation = ""

    for line in lines[:15]:
        lowered = line.lower()

        if (
            len(line) <= 80
            and "@" not in line
            and not line[0].islower()
            and any(word in lowered for word in DESIGNATION_WORDS)
            and not lowered.startswith(("career", "profile", "summary"))
        ):
            designation = line.strip(" |-–•")
            break

    return {
        "phone": phone,
        "linkedin": linkedin,
        "experience": experience,
        "location": location,
        "qualifications": ", ".join(qualifications),
        "skills": ", ".join(skills),
        "designation": designation,
    }


def export_structured_excel():
    """
    Write a clean, one-row-per-candidate sheet to
    candidates_structured.xlsx, built from each candidate.json
    and its parsed resume text.
    """

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = []

    for candidate_dir in sorted(CANDIDATES_DIR.iterdir()):

        json_file = candidate_dir / "candidate.json"

        if not candidate_dir.is_dir() or not json_file.exists():
            continue

        try:
            data = json.loads(json_file.read_text(encoding="utf-8"))
        except Exception as e:
            log(f"Could not read {json_file}: {e}", "WARNING")
            continue

        parsed = structured_fields_from_resume(data.get("resume_text", ""))

        rows.append({
            "Name": data.get("name") or data.get("parsed_name", ""),
            "Email": data.get("parsed_email", ""),
            "Phone": parsed["phone"] or data.get("parsed_phone", ""),
            "Current Designation": (
                data.get("current_designation") or parsed["designation"]
            ),
            "Current Company": data.get("current_company", ""),
            "Location": data.get("location") or parsed["location"],
            "Total Experience (Years)": (
                data.get("total_experience_years") or parsed["experience"]
            ),
            "Notice Period": data.get("notice_period", ""),
            "Qualifications": parsed["qualifications"],
            "Key Skills": parsed["skills"],
            "LinkedIn": parsed["linkedin"],
            "Resume File": str(candidate_dir / data.get("resume_file", "")),
            "Naukri Profile": data.get("profile_url", ""),
            "Search URL": data.get("search_url", ""),
            "Collected At": data.get("collected_at", ""),
            "Downloaded At": data.get("downloaded_at", ""),
            "Status": data.get("status", ""),
            "Candidate ID": data.get("candidate_id", ""),
        })

    if not rows:
        log("No candidates available for structured Excel export.")
        return

    df = pd.DataFrame(rows)
    df.insert(0, "S.No", range(1, len(df) + 1))

    try:
        with pd.ExcelWriter(STRUCTURED_EXCEL_FILE, engine="openpyxl") as writer:

            df.to_excel(writer, index=False, sheet_name="Candidates")
            sheet = writer.sheets["Candidates"]

            header_fill = PatternFill("solid", fgColor="1F4E78")

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            widths = {
                "S.No": 6, "Name": 26, "Email": 32, "Phone": 16,
                "Current Designation": 34, "Current Company": 24,
                "Location": 16, "Total Experience (Years)": 14,
                "Notice Period": 14, "Qualifications": 24,
                "Key Skills": 50, "LinkedIn": 30, "Resume File": 30,
                "Naukri Profile": 18, "Search URL": 18,
                "Collected At": 20, "Downloaded At": 20,
                "Status": 10, "Candidate ID": 18,
            }

            link_columns = {"LinkedIn", "Resume File", "Naukri Profile"}

            for index, column in enumerate(df.columns, start=1):
                letter = get_column_letter(index)
                sheet.column_dimensions[letter].width = widths.get(column, 18)

                for cell in sheet[letter][1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

                    if column in link_columns and cell.value:
                        target = str(cell.value)

                        if column == "LinkedIn" and not target.startswith("http"):
                            target = "https://" + target

                        cell.hyperlink = target
                        cell.font = Font(color="0563C1", underline="single")

                        if column == "Naukri Profile":
                            cell.value = "Open profile"
                        elif column == "Resume File":
                            cell.value = Path(target).name

            sheet.freeze_panes = "C2"
            sheet.auto_filter.ref = sheet.dimensions

    except PermissionError:
        log(
            f"Could not write {STRUCTURED_EXCEL_FILE} - close it in "
            "Excel and run with --export.",
            "WARNING",
        )
        return

    log(
        f"Structured Excel updated: {STRUCTURED_EXCEL_FILE}",
        "SUCCESS",
    )


# ============================================================
# CHROME / CDP
# ============================================================

def find_chrome_executable():
    """Find Google Chrome on Windows."""

    candidates = [
        Path(os.environ.get("LOCALAPPDATA", ""))
        / "Google" / "Chrome" / "Application" / "chrome.exe",

        Path(os.environ.get("PROGRAMFILES", ""))
        / "Google" / "Chrome" / "Application" / "chrome.exe",

        Path(os.environ.get("PROGRAMFILES(X86)", ""))
        / "Google" / "Chrome" / "Application" / "chrome.exe",
    ]

    for candidate in candidates:

        if candidate.exists():
            return candidate

    raise FileNotFoundError(
        "Google Chrome was not found."
    )


def cdp_browser_name():
    """
    Return the "Browser" field of the browser listening on the CDP port
    (e.g. "Chrome/141.0..." or "Edg/141.0..."), or None if nothing is.
    """

    try:
        import urllib.request

        with urllib.request.urlopen(
            f"{CDP_URL}/json/version",
            timeout=2,
        ) as response:

            return json.loads(response.read()).get("Browser", "")

    except Exception:
        return None


def cdp_available():
    """Check whether the automation Chrome is listening on port 9222."""

    name = cdp_browser_name()

    return bool(name) and name.startswith("Chrome/")


def is_chrome_running():
    """Return True if any chrome.exe process exists."""
    try:
        result = subprocess.run(
            ["tasklist", "/FI", "IMAGENAME eq chrome.exe", "/NH"],
            capture_output=True,
            text=True,
            timeout=5,
        )
        return "chrome.exe" in result.stdout
    except Exception:
        return False


def get_chrome_policy():
    """Return the current per-user Chrome UserDataDir policy, or None."""

    import winreg

    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, CHROME_POLICY_KEY) as key:
            value, _ = winreg.QueryValueEx(key, "UserDataDir")
            return value
    except OSError:
        return None


def ensure_chrome_policy():
    """
    Point Chrome's default profile directory at CHROME_PROFILE_DIR so the
    Naukri Launcher's `start chrome "<url>"` opens inside the automation
    Chrome. Returns True if the policy was changed just now.
    """

    import winreg

    wanted = str(CHROME_PROFILE_DIR)

    if get_chrome_policy() == wanted:
        return False

    try:
        with winreg.CreateKey(winreg.HKEY_CURRENT_USER, CHROME_POLICY_KEY) as key:
            winreg.SetValueEx(key, "UserDataDir", 0, winreg.REG_SZ, wanted)
    except PermissionError:
        # Windows only lets administrators write under Software\Policies.
        # An elevated prompt may run as a different (admin) account, so
        # address this user's registry hive by SID rather than HKCU.
        sid = subprocess.run(
            ["whoami", "/user", "/fo", "csv", "/nh"],
            capture_output=True,
            text=True,
        ).stdout.strip().split(",")[-1].strip('"')

        raise RuntimeError(
            "Setting Chrome's profile folder needs administrator rights "
            "(one time only).\n"
            "Open PowerShell as Administrator (Start menu > type "
            "PowerShell > Run as administrator) and run:\n\n"
            f'  reg add "HKU\\{sid}\\{CHROME_POLICY_KEY}" /v UserDataDir '
            f'/t REG_SZ /d "{wanted}" /f\n\n'
            "Then run this script again."
        ) from None

    log(
        "Chrome is now set to use the automation profile "
        f"({wanted}). Run `python main.py --restore-chrome` to undo.",
        "SUCCESS",
    )

    return True


def restore_chrome_policy():
    """Remove the UserDataDir policy so Chrome uses its normal profile."""

    import winreg

    if get_chrome_policy() is None:
        print("Chrome is already using its normal profile.")
        return

    try:
        with winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            CHROME_POLICY_KEY,
            0,
            winreg.KEY_SET_VALUE,
        ) as key:
            winreg.DeleteValue(key, "UserDataDir")
    except PermissionError:
        print("This needs administrator rights. In PowerShell opened with")
        print("'Run as administrator', run:")
        print()
        sid = subprocess.run(
            ["whoami", "/user", "/fo", "csv", "/nh"],
            capture_output=True,
            text=True,
        ).stdout.strip().split(",")[-1].strip('"')
        print(f'  reg delete "HKU\\{sid}\\{CHROME_POLICY_KEY}" /v UserDataDir /f')
        return

    print(
        "Chrome will use its normal profile again. "
        "Close every Chrome window for this to take effect."
    )


def start_chrome_debug():
    """
    Start the automation Chrome with CDP on port 9222.

    Chrome refuses remote debugging on its built-in default profile
    directory, so the profile is passed explicitly with --user-data-dir.
    The UserDataDir policy points at the same folder, so any later
    `chrome <url>` (from the Naukri Launcher) is handed to this running
    instance and opens as a tab we can see.
    """

    ensure_chrome_policy()

    other = cdp_browser_name()

    if other is not None and not other.startswith("Chrome/"):

        # Left over from older runs that automated Microsoft Edge.
        print()
        print(f"Another browser ({other}) is using automation port 9222.")
        print("This is the old automation Edge window from earlier runs.")
        print("Close that Edge window, then press ENTER to continue.")
        input()

        deadline = time.time() + 15

        while cdp_browser_name() is not None and time.time() < deadline:
            time.sleep(0.5)

        if cdp_browser_name() is not None:
            raise RuntimeError(
                "Port 9222 is still in use. Close the old automation "
                "Edge window and run the script again."
            )

    if cdp_available():

        log(
            "Chrome CDP is already available.",
            "SUCCESS",
        )

        return

    if is_chrome_running():

        # A Chrome started without the debugging port owns the profile,
        # so a new launch would just open a window in it without CDP.
        print()
        print("Google Chrome is already open without automation enabled.")
        print("Close ALL Chrome windows, then press ENTER to continue.")
        input()

        deadline = time.time() + 15

        while is_chrome_running() and time.time() < deadline:
            time.sleep(0.5)

        if is_chrome_running():
            raise RuntimeError(
                "Chrome is still running. Close every Chrome window "
                "(check the system tray too) and run the script again."
            )

    chrome = find_chrome_executable()

    log(f"Starting automation Chrome: {chrome}")

    subprocess.Popen(
        [
            str(chrome),
            "--remote-debugging-port=9222",
            f"--user-data-dir={CHROME_PROFILE_DIR}",
            "--no-first-run",
            "--no-default-browser-check",
            "https://recruit.naukri.com/",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=getattr(
            subprocess,
            "CREATE_NEW_PROCESS_GROUP",
            0,
        ),
    )

    deadline = time.time() + 30

    while time.time() < deadline:

        if cdp_available():

            log("Chrome CDP is ready.", "SUCCESS")

            return

        time.sleep(0.5)

    raise RuntimeError(
        "Chrome started but CDP port 9222 did not become "
        "available within 30 seconds.\n"
        "This can happen if another process is using port 9222.\n"
        "Run: netstat -an | findstr 9222  to check."
    )



# ============================================================
# PLAYWRIGHT PAGE HELPERS
# ============================================================

def is_login_url(url):
    url = (url or "").lower()

    return (
        "/login" in url
        or "login?" in url
        or "msg=to" in url
        or "autologinfailure" in url
        or "errorcode=" in url
    )


def get_fresh_page(browser_or_context, playwright=None):
    """Return a live, unclosed Page object, reconnecting CDP if necessary."""
    browser = browser_or_context
    if hasattr(browser_or_context, "is_connected") and not browser_or_context.is_connected():
        log("Browser disconnected. Reconnecting to CDP...", "WARNING")
        start_chrome_debug()
        if playwright:
            browser = playwright.chromium.connect_over_cdp(CDP_URL)

    if hasattr(browser, "contexts"):
        contexts = browser.contexts
    else:
        contexts = [browser]

    for ctx in contexts:
        for p in ctx.pages:
            try:
                if not p.is_closed():
                    _ = p.url
                    return p, browser
            except Exception:
                continue

    try:
        if contexts:
            return contexts[0].new_page(), browser
    except Exception:
        pass

    if hasattr(browser, "new_context"):
        ctx = browser.new_context()
        return ctx.new_page(), browser

    raise RuntimeError("Could not find or create an active browser page.")



def get_usable_pages(context):

    pages = []

    for page in context.pages:

        try:
            url = page.url

            if not url:
                continue

            if url.startswith("chrome://"):
                continue

            if url.startswith("edge://"):
                continue

            if url.startswith("about:"):
                continue

            pages.append(page)

        except Exception:
            continue

    return pages


def get_logged_in_naukri_page(browser_or_context):

    """
    Find an already authenticated Naukri page across ALL browser contexts.

    Accepts either a Playwright Browser object (preferred — scans every
    context/window) or a single BrowserContext.

    Priority:
    1. Naukri Recruiter page that is not a login page.
    2. Other Naukri page that is not a login page.

    Scanning all contexts is essential when the Naukri Recruiter Launcher
    opens the dashboard in a new window: Playwright may place that window
    in a separate context from the one we originally created.
    """

    # Build the list of contexts to inspect.
    if hasattr(browser_or_context, "contexts"):
        # A Browser object — scan every context.
        contexts = browser_or_context.contexts
    else:
        # A single BrowserContext (backward-compat).
        contexts = [browser_or_context]

    def _all_pages():
        for ctx in contexts:
            yield from get_usable_pages(ctx)

    # First priority: Recruiter/Resdex/Enterprise pages.
    for page in _all_pages():

        try:
            url = page.url.lower()

            if not is_naukri_url(url):
                continue

            if is_login_url(url):
                continue

            if (
                "recruit.naukri.com" in url
                or "resdex.naukri.com" in url
                or "enterprise.naukri.com" in url
            ):
                return page

        except Exception:
            continue

    # Second priority: any authenticated Naukri page.
    for page in _all_pages():

        try:
            url = page.url.lower()

            if not is_naukri_url(url):
                continue

            if is_login_url(url):
                continue

            return page

        except Exception:
            continue

    return None


def start_browser(playwright):

    """
    Connect to the existing automation Chrome.

    If Chrome is not running with CDP, start the dedicated
    automation Chrome profile.
    """

    log("Connecting to Google Chrome...")

    start_chrome_debug()

    browser = playwright.chromium.connect_over_cdp(
        CDP_URL
    )

    if not browser.contexts:
        raise RuntimeError(
            "No Chrome browser context was found."
        )

    context = browser.contexts[0]

    pages = get_usable_pages(context)

    log(
        f"Connected to Chrome. Open pages: {len(pages)}",
        "SUCCESS",
    )

    for index, page in enumerate(
        pages,
        start=1
    ):

        try:
            log(
                f"Page {index}: {page.url}"
            )
        except Exception:
            pass

    logged_page = get_logged_in_naukri_page(
        browser
    )

    if logged_page:

        log(
            f"Existing logged-in Naukri page found: "
            f"{logged_page.url}",
            "SUCCESS",
        )

        return browser, context, logged_page

    # No logged-in page yet.
    if pages:

        page = pages[0]

    else:

        page = context.new_page()

    return browser, context, page


# ============================================================
# LOGIN / SESSION
# ============================================================

def is_naukri_logged_in(browser):

    page = get_logged_in_naukri_page(
        browser
    )

    return page is not None


def reconnect_if_needed(browser, playwright=None):
    """
    Return a connected Browser, restarting the automation Chrome if it
    was closed.

    browser.is_connected() can still report True for a moment after
    Chrome exits, so the CDP port itself is checked as well.
    """

    alive = (
        cdp_available()
        and hasattr(browser, "is_connected")
        and browser.is_connected()
    )

    if alive:
        return browser

    log(
        "Automation Chrome was closed. Restarting it and reconnecting...",
        "WARNING",
    )

    start_chrome_debug()

    if playwright:
        browser = playwright.chromium.connect_over_cdp(CDP_URL)

    return browser


def ensure_logged_in(page, browser, playwright=None):

    # --------------------------------------------------------
    # EXISTING SESSION
    # --------------------------------------------------------

    try:
        if hasattr(browser, "is_connected") and browser.is_connected():
            logged_page = get_logged_in_naukri_page(browser)
            if logged_page:
                log(
                    "Existing Naukri login detected.",
                    "SUCCESS",
                )
                return logged_page
    except Exception:
        pass

    # --------------------------------------------------------
    # NO SESSION — log in through the Naukri Launcher
    # --------------------------------------------------------

    log(
        "No authenticated Naukri session detected."
    )

    print()
    print("=" * 65)
    print("              NAUKRI LOGIN REQUIRED")
    print("=" * 65)
    print()
    print("  1. Open the Naukri Launcher.")
    print("  2. Choose GOOGLE CHROME as the browser.")
    print("  3. Log in with your Naukri Recruiter credentials.")
    print()
    print("  The Launcher will open the dashboard as a new tab in the")
    print("  automation Chrome window, and the script will continue")
    print("  on its own. There is nothing to paste here.")
    print()
    print(f"  Waiting up to {LOGIN_WAIT_SECONDS // 60} minutes. "
          "Press Ctrl+C to cancel.")
    print("=" * 65)
    print()

    deadline = time.time() + LOGIN_WAIT_SECONDS
    next_notice = time.time() + 30

    while time.time() < deadline:

        try:
            browser = reconnect_if_needed(browser, playwright)

            logged_page = get_logged_in_naukri_page(browser)

            if logged_page:
                logged_page.bring_to_front()
                log(
                    f"Naukri authenticated session detected: "
                    f"{logged_page.url}",
                    "SUCCESS",
                )
                return logged_page

        except Exception as err:
            log(f"Still waiting for login ({err})", "WARNING")

        if time.time() >= next_notice:
            log("Still waiting for the Launcher login...")
            next_notice = time.time() + 30

        time.sleep(2)

    raise RuntimeError(
        "No logged-in Naukri tab appeared in the automation Chrome.\n"
        "Make sure the Naukri Launcher is set to open Google Chrome. If the "
        "dashboard opened in a different Chrome window, close all Chrome "
        "windows and run the script again."
    )



# ============================================================
# USER INPUT
# ============================================================

def ask_naukri_base_url():

    print()
    print("=" * 65)
    print("              NAUKRI RECRUITER URL")
    print("=" * 65)
    print()
    print(
        "Enter the Naukri Recruiter dashboard URL."
    )
    print(
        "Example: https://recruit.naukri.com/"
    )
    print()
    print(
        "NOTE: Do NOT paste a login/SSO/token URL."
    )
    print(
        "      Use the URL of your logged-in dashboard."
    )
    print()

    while True:

        url = input(
            "Naukri URL: "
        ).strip()

        if not url:
            print(
                "Please enter a URL."
            )
            continue

        if not url.startswith(
            ("http://", "https://")
        ):
            print(
                "Please enter the complete URL including https://"
            )
            continue

        if not is_naukri_url(url):

            print(
                "That does not appear to be a Naukri URL."
            )

            continue

        if is_login_url(url):

            print()
            print(
                "That looks like a login / SSO / error URL."
            )
            print(
                "Please enter your Recruiter dashboard URL instead."
            )
            print(
                "Example: https://recruit.naukri.com/"
            )
            print()

            continue

        return url



def ask_search_url():

    print()
    print("=" * 65)
    print("              CANDIDATE SEARCH URL")
    print("=" * 65)
    print()
    print(
        "Paste the authorized Naukri candidate-search URL."
    )
    print()

    while True:

        url = input(
            "Search URL: "
        ).strip()

        if not url:
            print(
                "Please enter a URL."
            )
            continue

        if not url.startswith(
            ("http://", "https://")
        ):
            print(
                "Please enter the complete URL including https://"
            )
            continue

        if not is_naukri_url(url):

            print(
                "That does not appear to be a Naukri URL."
            )

            continue

        return url


def get_page_number(url):
    """Return the pageNo value of a Resdex search URL (default 1)."""

    try:
        return int(parse_qs(urlparse(url).query).get("pageNo", ["1"])[0])
    except ValueError:
        return 1


def set_page_number(url, page_no):
    """Return the search URL with pageNo set to page_no."""

    parsed = urlparse(url)
    query = parse_qs(parsed.query, keep_blank_values=True)
    query["pageNo"] = [str(page_no)]

    return parsed._replace(
        query=urlencode(query, doseq=True)
    ).geturl()


# ============================================================
# OPEN RECRUITER BASE URL
# ============================================================

def open_base_url(page, base_url):

    log(
        f"Opening Naukri Recruiter URL: {base_url}"
    )

    current_url = page.url or ""

    # Avoid unnecessary reload if we're already there.
    if current_url.rstrip("/") != base_url.rstrip("/"):

        page.goto(
            base_url,
            wait_until="domcontentloaded",
        )

    time.sleep(2)

    log(
        f"Current Naukri page: {page.url}",
        "SUCCESS",
    )


# ============================================================
# CANDIDATE SEARCH
# ============================================================

def captcha_visible(page):
    """True if Naukri is showing its "Select all images..." check."""

    try:
        for frame in page.frames:
            if "recaptcha/api2/bframe" in (frame.url or ""):
                element = frame.frame_element()
                if element.is_visible():
                    return True

        return "Select all images" in page.locator("body").inner_text(
            timeout=5_000
        )

    except Exception:
        return False


def wait_if_captcha(page):
    """
    If Naukri shows a CAPTCHA (it does after many profiles are opened
    quickly), pause until the user solves it in Chrome. Returns True if
    a CAPTCHA was shown.
    """

    if not captcha_visible(page):
        return False

    log(
        "Naukri is showing a CAPTCHA (image check).",
        "WARNING",
    )

    print()
    print("=" * 65)
    print("  NAUKRI CAPTCHA")
    print("=" * 65)
    print("  1. Switch to the automation Chrome window.")
    print("  2. Solve the image check and click VERIFY.")
    print("  3. Come back here and press ENTER.")
    print("=" * 65)
    input("Press ENTER after solving the CAPTCHA: ")

    return True


def open_candidate_search(page, search_url):

    log(
        f"Opening candidate search URL: {search_url}"
    )

    page.goto(
        search_url,
        wait_until="domcontentloaded",
    )

    # Give client-side UI time to render.
    try:
        page.wait_for_load_state(
            "networkidle",
            timeout=20_000,
        )
    except Exception:
        pass

    time.sleep(2)

    if is_login_url(page.url):

        raise RuntimeError(
            "Naukri redirected to the login page while "
            "opening the candidate search URL. "
            "The current Naukri session may have expired."
        )

    if wait_if_captcha(page):

        # Results are only loaded after the check is passed.
        page.reload(wait_until="domcontentloaded")

        try:
            page.locator(".tuple-card").first.wait_for(timeout=20_000)
        except Exception:
            pass

    log(
        f"Candidate search opened: {page.url}",
        "SUCCESS",
    )


# ============================================================
# CANDIDATE COLLECTION
# ============================================================

def collect_candidates(page, limit):

    candidates = []

    log(
        f"CURRENT URL: {page.url}"
    )

    try:
        log(
            f"PAGE TITLE: {page.title()}"
        )

    except Exception:
        pass

    # --------------------------------------------------------
    # Diagnostics
    # --------------------------------------------------------

    try:

        body_text = page.locator(
            "body"
        ).inner_text(
            timeout=5_000
        )

        log(
            f"PAGE TEXT: {body_text[:1000]}"
        )

    except Exception as e:

        log(
            f"Could not read page text: {e}",
            "WARNING",
        )

    # --------------------------------------------------------
    # Candidate cards
    # --------------------------------------------------------

    cards = page.locator(
        ".tuple-card"
    )

    total = cards.count()

    log(
        f"Candidate cards detected: {total}"
    )

    if total == 0:

        log(
            "No .tuple-card elements were found.",
            "WARNING",
        )

        return candidates

    maximum = min(
        total,
        limit
    )

    for index in range(maximum):

        try:

            card = cards.nth(index)

            # ------------------------------------------------
            # Candidate name
            # ------------------------------------------------

            name = ""

            try:

                links = card.locator("a")

                if links.count() > 0:

                    name = (
                        links.nth(0)
                        .inner_text(
                            timeout=3_000
                        )
                        .strip()
                    )

            except Exception:
                pass

            if not name:

                try:
                    name = card.inner_text(
                        timeout=3_000
                    ).splitlines()[0].strip()

                except Exception:
                    name = f"candidate_{index + 1}"

            # ------------------------------------------------
            # Candidate profile URL
            # ------------------------------------------------

            profile_url = ""

            try:

                links = card.locator("a")

                for link_index in range(
                    min(links.count(), 10)
                ):

                    href = links.nth(
                        link_index
                    ).get_attribute("href")

                    if not href:
                        continue

                    if href.startswith("/"):

                        parsed = urlparse(
                            page.url
                        )

                        href = (
                            f"{parsed.scheme}://"
                            f"{parsed.netloc}"
                            f"{href}"
                        )

                    if href.startswith(
                        ("http://", "https://")
                    ):

                        if "naukri.com" in href.lower():

                            profile_url = href
                            break

            except Exception:
                pass

            # ------------------------------------------------
            # Candidate ID
            # ------------------------------------------------

            candidate_id = ""

            if profile_url:

                # Resdex profile links look like
                # /v3/preview?...&uniqId=<id>&uresid=<id>...
                query = parse_qs(urlparse(profile_url).query)

                for key in ("uniqId", "uresid"):
                    if query.get(key):
                        candidate_id = query[key][0]
                        break

                if not candidate_id:

                    match = re.search(
                        r"(?:[?&]id=|profile/|candidate/)([A-Za-z0-9_-]+)",
                        profile_url,
                        re.IGNORECASE,
                    )

                    if match:
                        candidate_id = match.group(1)

            if not candidate_id:

                candidate_id = f"candidate_{index + 1}"

            candidate = {
                "candidate_id": candidate_id,
                "name": name,
                "profile_url": profile_url,
                "search_url": page.url,
                "collected_at": datetime.now().isoformat(),
            }

            candidates.append(candidate)

            log(
                f"Candidate {index + 1}: "
                f"{name} | {profile_url}"
            )

        except Exception as e:

            log(
                f"Could not collect candidate "
                f"{index + 1}: {e}",
                "WARNING",
            )

    return candidates


# ============================================================
# CANDIDATE PROFILE
# ============================================================

def process_candidate(page, candidate):

    log(
        f"Processing: {candidate.get('name', 'Unknown')}"
    )

    profile_url = candidate.get(
        "profile_url"
    )

    if not profile_url:

        raise RuntimeError(
            "Candidate profile URL was not found."
        )

    # --------------------------------------------------------
    # Open candidate profile
    # --------------------------------------------------------

    page.goto(
        profile_url,
        wait_until="domcontentloaded",
    )

    try:

        page.wait_for_load_state(
            "networkidle",
            timeout=20_000,
        )

    except Exception:
        pass

    # Opening profiles back-to-back triggers Naukri's CAPTCHA, so pace
    # the requests a little like a person would.
    time.sleep(random.uniform(*PROFILE_DELAY_SECONDS))

    if is_login_url(page.url):

        raise RuntimeError(
            "Candidate profile redirected to the Naukri login page."
        )

    if wait_if_captcha(page):

        page.reload(wait_until="domcontentloaded")

        try:
            page.wait_for_load_state("networkidle", timeout=20_000)
        except Exception:
            pass

    # --------------------------------------------------------
    # Candidate fields
    #
    # These can be populated after we confirm the actual
    # selectors from your Naukri Recruiter UI.
    # --------------------------------------------------------

    candidate["location"] = ""
    candidate["current_company"] = ""
    candidate["current_designation"] = ""
    candidate["total_experience_years"] = None
    candidate["notice_period"] = ""

    # --------------------------------------------------------
    # RESUME DOWNLOAD
    # --------------------------------------------------------

    download_button = page.locator(
        'button:has-text("Download CV")'
    )

    button_count = download_button.count()

    if button_count == 0:

        # Some UIs may render the control as a link.
        download_button = page.locator(
            'a:has-text("Download CV")'
        )

        button_count = download_button.count()

    if button_count == 0:

        raise RuntimeError(
            "Permitted 'Download CV' control was not found."
        )

    # --------------------------------------------------------
    # Download
    # --------------------------------------------------------

    with page.expect_download(
        timeout=60_000
    ) as download_info:

        download_button.first.click()

    download = download_info.value

    filename = safe_filename(
        download.suggested_filename
    )

    if not filename:

        filename = (
            f"{safe_filename(candidate.get('name'))}.pdf"
        )

    temporary_file = (
        DATA_DIR / filename
    )

    # Prevent accidental overwrite.
    if temporary_file.exists():

        stem = temporary_file.stem
        suffix = temporary_file.suffix

        temporary_file = (
            DATA_DIR
            / f"{stem}_{int(time.time())}{suffix}"
        )

    download.save_as(
        str(temporary_file)
    )

    if not temporary_file.exists():

        raise RuntimeError(
            "Resume download failed."
        )

    if temporary_file.stat().st_size == 0:

        raise RuntimeError(
            "Downloaded resume is empty."
        )

    log(
        f"Resume downloaded: {temporary_file.name}",
        "SUCCESS",
    )

    return temporary_file


# ============================================================
# PROCESS CANDIDATES
# ============================================================

def process_candidates(
    page,
    candidates,
    progress,
    limit,
):

    processed_this_run = 0

    for candidate in candidates:

        if processed_this_run >= limit:
            break

        try:

            log(
                f"--------------------------------------------------"
            )

            log(
                f"Processing "
                f"{processed_this_run + 1}/{limit}: "
                f"{candidate.get('name')}"
            )

            # ----------------------------------------------
            # Skip candidates saved in an earlier run, so
            # their profile is not opened (and CV quota is
            # not used) again.
            # ----------------------------------------------

            saved_prefix = (
                safe_filename(candidate.get("candidate_id", "")) + "_"
            )

            if not candidate.get("candidate_id", "").startswith("candidate_") and any(
                folder.name.startswith(saved_prefix)
                and (folder / "candidate.json").exists()
                for folder in CANDIDATES_DIR.iterdir()
            ):

                log(
                    f"Already downloaded earlier, skipping: "
                    f"{candidate.get('name')}"
                )

                continue

            # ----------------------------------------------
            # Download
            # ----------------------------------------------

            resume_path = process_candidate(
                page,
                candidate,
            )

            # ----------------------------------------------
            # Hash
            # ----------------------------------------------

            resume_hash = sha256_file(
                resume_path
            )

            # ----------------------------------------------
            # Duplicate
            # ----------------------------------------------

            duplicate = find_duplicate(
                resume_hash
            )

            if duplicate:

                log(
                    f"Duplicate resume detected for "
                    f"{candidate.get('name')}",
                    "DUPLICATE",
                )

                resume_path.unlink(
                    missing_ok=True
                )

                progress["duplicates"] += 1

                # A duplicate was checked but not saved
                # as a new candidate.
                save_progress(progress)

                processed_this_run += 1

                continue

            # ----------------------------------------------
            # Save + parse
            # ----------------------------------------------

            saved_candidate = save_candidate(
                candidate,
                resume_path,
            )

            progress["successful"] += 1
            progress["processed"] += 1

            save_progress(progress)

            log(
                f"Saved candidate: "
                f"{saved_candidate.get('name')}",
                "SUCCESS",
            )

            processed_this_run += 1

        except Exception as e:

            progress["failed"] += 1
            progress["processed"] += 1

            save_progress(progress)

            log(
                f"Failed "
                f"{candidate.get('name')}: {e}",
                "ERROR",
            )

            failed_id = safe_filename(
                candidate.get(
                    "candidate_id",
                    "unknown",
                )
            )

            failed_folder = (
                FAILED_DIR / failed_id
            )

            failed_folder.mkdir(
                parents=True,
                exist_ok=True,
            )

            try:

                (
                    failed_folder
                    / "candidate.json"
                ).write_text(
                    json.dumps(
                        candidate,
                        indent=2,
                        ensure_ascii=False,
                    ),
                    encoding="utf-8",
                )

            except Exception as save_error:

                log(
                    f"Could not save failed candidate: "
                    f"{save_error}",
                    "WARNING",
                )

            processed_this_run += 1

    return processed_this_run


# ============================================================
# MAIN RUN
# ============================================================

def run(limit):

    progress = load_progress()

    log(
        "=================================================="
    )

    log(
        "Starting Naukri Recruiter automation"
    )

    log(
        f"Requested candidate limit: {limit}"
    )

    # --------------------------------------------------------
    # Ask for the Naukri base URL first.
    # --------------------------------------------------------

    base_url = ask_naukri_base_url()

    # --------------------------------------------------------
    # Connect to existing automation Chrome.
    # --------------------------------------------------------

    with sync_playwright() as playwright:

        browser = None

        try:

            browser, context, page = start_browser(
                playwright
            )

            # ------------------------------------------------
            # Open the requested Naukri base URL.
            # ------------------------------------------------

            open_base_url(
                page,
                base_url,
            )

            # ------------------------------------------------
            # Check login.
            # ------------------------------------------------

            page = ensure_logged_in(
                page,
                browser,
                playwright,
            )

            # ------------------------------------------------
            # Ask for search URL.
            # ------------------------------------------------

            search_url = ask_search_url()

            # ------------------------------------------------
            # Go through result pages until the limit is
            # reached or the results run out.
            # ------------------------------------------------

            processed = 0
            page_no = get_page_number(search_url)

            while processed < limit:

                page_url = set_page_number(search_url, page_no)

                log(f"========== Results page {page_no} ==========")

                open_candidate_search(
                    page,
                    page_url,
                )

                candidates = collect_candidates(
                    page,
                    limit - processed,
                )

                log(
                    f"Found {len(candidates)} candidates on page {page_no}.",
                    "SUCCESS",
                )

                if not candidates:

                    log(
                        "No more candidates found. Stopping.",
                        "WARNING",
                    )

                    break

                processed += process_candidates(
                    page,
                    candidates,
                    progress,
                    limit - processed,
                )

                page_no += 1

            log(
                f"Processed this run: {processed}",
                "SUCCESS",
            )

            # ------------------------------------------------
            # Export.
            # ------------------------------------------------

            export_excel()

            log(
                "Processing completed.",
                "SUCCESS",
            )

            log(
                f"Successful: {progress.get('successful', 0)}"
            )

            log(
                f"Failed: {progress.get('failed', 0)}"
            )

            log(
                f"Duplicates: {progress.get('duplicates', 0)}"
            )

        except KeyboardInterrupt:

            log(
                "Automation stopped by user.",
                "WARNING",
            )

        except Exception as e:

            log(
                f"Automation stopped: {e}",
                "ERROR",
            )

            raise

        finally:

            # IMPORTANT:
            # Do NOT call browser.close().
            #
            # The Chrome window belongs to the user and should
            # remain open after the script finishes.
            #
            # Playwright will disconnect when the context
            # exits; the remote Chrome process remains running.

            pass


# ============================================================
# CLI
# ============================================================

if __name__ == "__main__":

    parser = argparse.ArgumentParser(
        description=(
            "Naukri Recruiter resume automation"
        )
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=1,
        help=(
            "Number of candidates to process "
            "(default: 1)"
        ),
    )

    parser.add_argument(
        "--export",
        action="store_true",
        help=(
            "Export existing candidates to Excel"
        ),
    )

    parser.add_argument(
        "--status",
        action="store_true",
        help=(
            "Show current progress"
        ),
    )

    parser.add_argument(
        "--restore-chrome",
        action="store_true",
        help=(
            "Make Chrome use its normal profile again "
            "(undoes the automation profile setting)"
        ),
    )

    args = parser.parse_args()

    if args.restore_chrome:

        restore_chrome_policy()

    elif args.status:

        progress = load_progress()

        print(
            json.dumps(
                progress,
                indent=2,
            )
        )

    elif args.export:

        export_excel()

    else:

        requested_limit = max(
            1,
            args.limit,
        )

        run(
            min(
                requested_limit,
                DAILY_LIMIT,
            )
        )